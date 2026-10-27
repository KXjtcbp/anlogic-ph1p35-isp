# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document

P = r"E:\FPGA\项目学习清单_20261023.docx"
d = Document(P)
print("段数 %d  表格 %d" % (len(d.paragraphs), len(d.tables)))
print()
print("=== 一级标题（>=15pt）===")
for p in d.paragraphs:
    if p.runs and p.runs[0].font.size and p.runs[0].font.size.pt >= 15:
        print("  " + p.text)
print()
print("=== 代码块（▎ 标注行）===")
n = 0
for p in d.paragraphs:
    t = p.text.strip()
    if t.startswith("\u258e"):
        n += 1
        print("  %2d. %s" % (n, t))
print("  共 %d 个代码块；▸ 解释行 %d 条" % (n, sum(1 for p in d.paragraphs if p.text.strip().startswith("\u25b8"))))
print()
print("=== 表格概览 ===")
for i, t in enumerate(d.tables):
    if t.rows:
        cells = [c.text.strip().replace("\n", " ")[:12] for c in t.rows[0].cells]
        print("  T%02d %2d行×%d列  %s" % (i + 1, len(t.rows), len(t.columns), " | ".join(cells)))
