# -*- coding: utf-8 -*-
"""第三步补丁：让 inline() 支持 italic，并让 rich() 的 (text, kw) 元组分支也走 inline()。
   然后重新装配、生成、复审。"""
import os, sys, re, importlib.util, hashlib
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
T = os.environ.get("TEMP", r"C:\Users\KX\AppData\Local\Temp")
L0 = os.path.join(T, "_L0.py")
s = open(L0, encoding="utf-8-sig").read()

# 1) inline() 加 italic
if "def inline(p, text, size=BODY, color=None, bold=False, mono=False, italic=False)" in s:
    print("[1] inline 已支持 italic，跳过")
else:
    old_sig = "def inline(p, text, size=BODY, color=None, bold=False, mono=False):"
    new_sig = "def inline(p, text, size=BODY, color=None, bold=False, mono=False, italic=False):"
    assert old_sig in s, "找不到 inline 签名"
    s = s.replace(old_sig, new_sig, 1)
    # 三处 run(...) 调用补 italic=italic
    s = s.replace("run(p, seg[2:-2], bold=True, size=size, color=color, mono=mono)",
                  "run(p, seg[2:-2], bold=True, size=size, color=color, mono=mono, italic=italic)", 1)
    s = s.replace("run(p, seg[1:-1], bold=bold, size=size, color=color, mono=True)",
                  "run(p, seg[1:-1], bold=bold, size=size, color=color, mono=True, italic=italic)", 1)
    s = s.replace("run(p, seg, bold=bold, size=size, color=color, mono=mono)",
                  "run(p, seg, bold=bold, size=size, color=color, mono=mono, italic=italic)", 1)
    print("[1] inline 已加 italic")

# 2) rich() 元组分支改走 inline
old_rich = '''            txt, kw = t
            run(p, txt, bold=kw.get("b", False), size=kw.get("s", size),
                color=kw.get("c"), mono=kw.get("m", False), italic=kw.get("i", False))'''
new_rich = '''            txt, kw = t
            inline(p, txt, size=kw.get("s", size), color=kw.get("c"),
                   bold=kw.get("b", False), mono=kw.get("m", False),
                   italic=kw.get("i", False))'''
if new_rich in s:
    print("[2] rich 元组分支已改过，跳过")
else:
    assert old_rich in s, "找不到 rich 元组分支"
    s = s.replace(old_rich, new_rich, 1)
    print("[2] rich 元组分支已改走 inline")

open(L0, "w", encoding="utf-8").write(s)
print("    _L0.py 长度 %d" % len(s))

# 3) 装配 + 生成
parts = [open(os.path.join(T, nm), encoding="utf-8-sig").read().replace("\ufeff", "")
         for nm in ("_L0.py", "_L1.py", "_L2.py", "_L3.py", "_L4.py")]
ALL = "\n".join(parts)
pa = os.path.join(T, "_ALL.py")
open(pa, "w", encoding="utf-8").write(ALL)
print("[3] _ALL.py %d chars, BOM %d" % (len(ALL), ALL.count("\ufeff")))

spec = importlib.util.spec_from_file_location("gen", pa)
mod = importlib.util.module_from_spec(spec)
sys.argv = ["gen"]
spec.loader.exec_module(mod)

OUT = r"E:\FPGA\项目学习清单_20261023.docx"
from docx import Document
d = Document(OUT)
ps = [p.text for p in d.paragraphs]
print("\n=== 结果 ===")
print("  %d B  md5=%s  段落 %d  表格 %d"
      % (os.path.getsize(OUT), hashlib.md5(open(OUT, "rb").read()).hexdigest(),
         len(ps), len(d.tables)))

print("\n=== 复审 ===")
b1 = [(i, ps[i][:60]) for i, t in enumerate(ps) if "**" in t]
b2 = [(i, ps[i][:60]) for i, t in enumerate(ps) if "`" in t]
print("  残留 ** : %s" % (b1 if b1 else "无"))
print("  残留 `  : %s" % (b2 if b2 else "无"))
srcs = [t for t in ps if t.startswith("▎")]
notes = [t for t in ps if t.startswith("▸")]
print("  代码块 %d / ▸ 解释 %d" % (len(srcs), len(notes)))
dup = {}
for s2 in srcs:
    dup[s2] = dup.get(s2, 0) + 1
print("  重复块: %s" % ({k: v for k, v in dup.items() if v > 1} or "无"))
