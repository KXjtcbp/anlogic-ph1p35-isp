# -*- coding: utf-8 -*-
"""补丁 2：补 data128_96 的完整例化 + 把图 3 / 图 4 挂到真实标题下。"""
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"E:\FPGA\项目学习清单_20261023.docx"
FIG = r"E:\FPGA\_analysis\out\figs"


def mkrpr(bold=False, mono=False, size=10.5, color=None):
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    if mono:
        rf.set(qn("w:ascii"), "Consolas")
        rf.set(qn("w:hAnsi"), "Consolas")
        rf.set(qn("w:eastAsia"), "Consolas")
    else:
        rf.set(qn("w:ascii"), "Times New Roman")
        rf.set(qn("w:hAnsi"), "Times New Roman")
        rf.set(qn("w:eastAsia"), "宋体")
    rPr.append(rf)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    return rPr


def add_run_mixed(p, text, size=8.0):
    if "//" in text:
        code, _, cmt = text.partition("//")
        for seg, col in ((code, "1F3864"), ("//" + cmt, "2E7D32")):
            if not seg:
                continue
            r = OxmlElement("w:r")
            r.append(mkrpr(mono=True, size=size, color=col))
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = seg
            r.append(t)
            p._p.append(r)
    else:
        r = OxmlElement("w:r")
        r.append(mkrpr(mono=True, size=size, color="1F3864"))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p._p.append(r)


def new_para(text="", bold=False, size=10.5, color=None, indent=True,
             before=0, after=4, spacing=1.38, align=None):
    p = Document().add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:line"), str(int(240 * spacing)))
    sp.set(qn("w:lineRule"), "auto")
    sp.set(qn("w:before"), str(int(before * 20)))
    sp.set(qn("w:after"), str(int(after * 20)))
    pPr.append(sp)
    if indent:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLine"), str(int(size * 2 * 20)))
        pPr.append(ind)
    if align:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), align)
        pPr.append(jc)
    if text:
        r = OxmlElement("w:r")
        r.append(mkrpr(bold=bold, mono=False, size=size, color=color))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p._p.append(r)
    return p


def code_para(line, size=8.0):
    p = new_para(indent=False, before=0, after=0, spacing=1.0)
    pPr = p._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "200")
    pPr.append(ind)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F4F4")
    pPr.append(shd)
    add_run_mixed(p, line, size=size)
    return p


def pt(p):
    return "".join(n.text or "" for n in p._p.iter(qn("w:t")))


D128 = [
    "data128_96 u_data128_96 (                     // 位宽转换：128 bit/拍 → 96 bit/拍",
    "    .I_tdata(m_axis_tdata),                   // 进来：4 像素 × 32 bit（R/G/B 各 8 位 + 8 位填充）",
    "    .O_tdata(m_axis_tdata_96)                 // 出去：4 像素 × 24 bit（R/G/B 正好各 8，不浪费一位）",
    ");                                            // 它只有 2 个端口、没有 tvalid/tuser —— 因为只是重排，不改变节奏",
    "// ★为什么非要有这一级★ 去马赛克输出是 4×32bit（每个像素用 32 位存 24 位信息，为对齐方便），",
    "//   而后面所有模块都只要 24 位。不转的话每拍浪费 32 bit 的存储与查表口，资源立刻吃紧。",
]

FIGS = [
    ("fig3_isp.png", "3.1 先理解为什么是这六个", "图 3　ISP 六个环节：顺序不能随便换",
     "图里每个方块下面的小字就是「为什么必须在这个位置」，比背顺序有用。"),
    ("fig4_blocks_osd.png", "6.1 为什么是「块级」而不是「像素级」", "图 4　块级检测的「块」是怎么切的｜八层 OSD 的叠放顺序",
     "左边是运动检测的 144 个块怎么切；右边是八层 OSD 谁压谁。"),
]


def main():
    doc = Document(DOC)
    paras = doc.paragraphs

    # 1) data128_96：插在 demosaic 那段 ▸ 解释之后
    ins = None
    for i, p in enumerate(paras):
        if pt(p).startswith("▎") and "isp_top.v" in pt(p) and "187~209" in pt(p):
            for j in range(i + 1, min(i + 60, len(paras))):
                if pt(paras[j]).startswith("▸"):
                    ins = j
                    break
            break
    if ins is None:
        print("  !! 找不到 demosaic 的 ▸ 行")
    else:
        prev = paras[ins]._p
        # 先加一个空段 + 小标题说明
        tip = new_para("紧跟着的下一级是位宽转换（4 像素从 32 位/像素压到 24 位/像素）：",
                       indent=True, after=3)
        prev.addnext(tip._p)
        prev = tip._p
        mk = new_para("▎hdl_source\\isp\\isp_top.v  第 211~214 行", size=9,
                      color="7A4A00", indent=False, after=2)
        prev.addnext(mk._p)
        prev = mk._p
        for ln in D128:
            np_ = code_para(ln)
            prev.addnext(np_._p)
            prev = np_._p
        nt = new_para("▸ 这段在做什么：这是「位宽转换」，把 4 个像素从 32 位/像素压成 24 位/像素。"
                      "它只有输入数据和输出数据两个端口，没有 tvalid/tuser——因为像素个数和节奏都没变，"
                      "只是把每个像素里多余的 8 位填充位扔掉。看起来多余，其实省下的是后面 6 级模块的存储和查表口。",
                      size=9.5, color="444444", indent=False, after=8)
        prev.addnext(nt._p)
        print("data128_96 已补入，%d 行" % len(D128))

    # 2) 两张图
    paras = doc.paragraphs
    for fn, anchor, cap, lead in FIGS:
        path = os.path.join(FIG, fn)
        if not os.path.exists(path):
            print("  !! 缺图:", fn)
            continue
        idx = None
        for i, p in enumerate(paras):
            if pt(p).strip().startswith(anchor):
                idx = i
                break
        if idx is None:
            for i, p in enumerate(paras):
                if anchor[:8] in pt(p):
                    idx = i
                    break
        if idx is None:
            print("  !! 找不到锚点:", anchor)
            continue
        p_lead = new_para(lead, indent=True, after=4)
        p_pic = Document().add_paragraph()
        p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pic.paragraph_format.space_before = Pt(4)
        p_pic.paragraph_format.space_after = Pt(2)
        p_pic.add_run().add_picture(path, width=Cm(15.6))
        p_cap = new_para(cap, size=9, color="595959", indent=False, after=10, align="center")
        paras[idx]._p.addnext(p_cap._p)
        paras[idx]._p.addnext(p_pic._p)
        paras[idx]._p.addnext(p_lead._p)
        print("已插图:", fn)

    doc.save(DOC)
    print("完成 →", DOC, os.path.getsize(DOC), "B")


if __name__ == "__main__":
    main()
