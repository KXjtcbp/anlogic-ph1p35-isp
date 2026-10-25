# -*- coding: utf-8 -*-
"""修正 sync_report_20261019.py 留下的两处问题：
1) 摘要里「目标检测」被「运动目标检测」这个词骗过去，没真正补上；
2) §3.2.3 的「（5）比特流版本台账」段被目标检测文案整段覆盖，版本台账丢失，需补回为（6）。
"""
import copy
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

F = r"E:\FPGA\比赛报告_填空版_20261023.docx"
d = Document(F)
log = []


def find(prefix):
    for i, p in enumerate(d.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    return -1


def replace_text(i, new):
    p = d.paragraphs[i]
    if not p.runs:
        p.add_run(new)
    else:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""


def insert_after(i, text):
    src = d.paragraphs[i]
    new_p = copy.deepcopy(src._p)
    src._p.addnext(new_p)
    np = Paragraph(new_p, src._parent)
    if not np.runs:
        np.add_run("")
    np.runs[0].text = text
    for r in np.runs[1:]:
        r.text = ""
    return np


# ---- 1) 摘要补目标检测 ----
i = find("在扩展与进阶功能方面")
if i >= 0:
    old = d.paragraphs[i].text.strip()
    if "硬件目标检测" not in old:
        replace_text(i, old.rstrip("。") +
                     "。此外还实现了按颜色特征识别的硬件目标检测：块级判决 + 两帧确认去抖 + "
                     "孤立块滤除 + 1 秒跟踪保持，命中的目标以绿色框叠加，空场景零误检。")
        log.append(("摘要 补硬件目标检测", i))

# ---- 2) §3.2.3 补回比特流版本台账为 (6) ----
i = find("（5）目标检测截图")
if i >= 0:
    j = find("（6）比特流版本台账")
    if j < 0:
        insert_after(i, "（6）比特流版本台账：`proj/releases/` 下 27 个版本目录，每版含 `.bit`、压缩版、"
                        "`route.qor`、`phy.area` 与 `readMe.txt`（记录 md5、时序、资源）。")
        log.append(("§3.2.3 补回(6) 比特流版本台账", i))

d.save(F)
print("已保存:", F, os.path.getsize(F), "B")
for tag, idx in log:
    print("  · %-30s 段落 %s" % (tag, idx))

# ---- 复核：目标段落现状 ----
d2 = Document(F)
print("\n=== 复核 ===")
for key in ("在扩展与进阶功能方面", "（4）OSD 信息叠加", "（5）硬件运动目标检测", "（6）硬件目标检测",
            "（7）实时参数调节", "（8）自动曝光闭环", "① RTL 设计", "⑤ 版本管理",
            "（12）`obj_detect`", "（13）`hdmi_mixer`", "（14）`uicfgcs500_720p`",
            "（5）目标检测截图", "（6）比特流版本台账", "系统主要性能指标如下表",
            "硬件平台为安路 HX1P35A", "待补充——这一节需参赛队员本人撰写",
            "本部分按「整体", "（1）目标检测的进一步扩展", "（5）QSPI 固化"):
    k = find(key)
    print("  %-38s -> %s" % (key[:36], "OK 段落%d" % k if k >= 0 else "★缺失★"))
print("\n段数=%d 表数=%d" % (len(d2.paragraphs), len(d2.tables)))
print("\n=== 表1 ===")
for r in d2.tables[0].rows:
    print("   | " + " | ".join(c.text.strip()[:60] for c in r.cells))
print("\n=== 表2 ===")
for r in d2.tables[1].rows:
    print("   | " + " | ".join(c.text.strip()[:44] for c in r.cells))
