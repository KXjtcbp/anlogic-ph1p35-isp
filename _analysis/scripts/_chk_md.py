# -*- coding: utf-8 -*-
import sys, re
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document

P = r"E:\FPGA\项目学习清单_20261023.docx"
d = Document(P)

print("=== 正文里残留的 Markdown / 排版标记 ===")
pats = {
    "**加粗星号": r"\*\*",
    "反引号": r"`",
    "Markdown 标题": r"^#{1,6}\s",
    "Markdown 列表": r"^[-*+]\s",
    "表格竖线": r"\|",
    "双横线破折号": r"——",
}
total = 0
for name, pat in pats.items():
    hits = []
    for i, p in enumerate(d.paragraphs):
        if re.search(pat, p.text):
            hits.append((i, p.text.strip()[:70]))
    # 表格里也查
    for ti, t in enumerate(d.tables):
        for r in t.rows:
            for c in r.cells:
                if re.search(pat, c.text):
                    hits.append(("T%02d" % (ti + 1), c.text.strip()[:50]))
    if hits:
        total += len(hits)
        print("\n  【%s】命中 %d 处" % (name, len(hits)))
        for i, t in hits[:8]:
            print("     %s  %s" % (i, t))
        if len(hits) > 8:
            print("     …还有 %d 处" % (len(hits) - 8))
print("\n合计 %d 处" % total)
