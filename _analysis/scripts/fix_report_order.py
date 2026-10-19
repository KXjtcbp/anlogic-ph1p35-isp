# -*- coding: utf-8 -*-
"""收尾修正：
1) §2.3 的 (12) obj_detect 被插到了 (13) hdmi_mixer 之后，顺序错，需物理前移到 ae_ctrl 之后；
2) 表 2 的「资源占用」「时序收敛」两行仍是旧数（slice 77.31% / LUT6 71.40% / WNS +102ps），
   与表 1 已更新的 82.24% / 77.30% / +104ps 冲突。
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

F = r"E:\FPGA\比赛报告_填空版_20261019.docx"
d = Document(F)
log = []


def find(prefix):
    for i, p in enumerate(d.paragraphs):
        if p.text.strip().startswith(prefix):
            return i
    return -1


# ---- 1) 把 obj_detect 段前移到 hdmi_mixer 之前 ----
i_obj = find("（12）`obj_detect`")
i_hdmi = find("（13）`hdmi_mixer`")
if i_obj >= 0 and i_hdmi >= 0 and i_obj > i_hdmi:
    p_obj = d.paragraphs[i_obj]._p
    p_hdmi = d.paragraphs[i_hdmi]._p
    parent = p_obj.getparent()
    parent.remove(p_obj)
    p_hdmi.addprevious(p_obj)
    log.append(("(12) obj_detect 前移到 (13) hdmi_mixer 之前", "%d→%d" % (i_obj, i_hdmi)))

# ---- 2) 表 2 两行旧数 ----
t2 = d.tables[1]
for r in range(len(t2.rows)):
    k = t2.cell(r, 0).text.strip()

    def setc(row, col, text):
        c = t2.cell(row, col)
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    if k == "资源占用":
        setc(r, 2, "slice 82.24%；LUT6 77.30%；ERAM 47.22%；DSP 25.00%（当前交付版）")
        log.append(("表2 资源占用 77.31→82.24", r))
    elif k == "时序收敛":
        setc(r, 2, "Setup WNS +104ps / TNS 0 / 0 违例；Hold WNS +20ps / TNS 0 / 0 违例")
        log.append(("表2 时序 +102ps→+104ps", r))

d.save(F)
print("已保存:", F, os.path.getsize(F), "B")
for tag, v in log:
    print("  · %-40s %s" % (tag, v))

# ---- 终检 ----
d2 = Document(F)
print("\n=== §2.3 模块清单顺序复核 ===")
for p in d2.paragraphs:
    t = p.text.strip()
    if len(t) > 3 and t[0] == "（" and t[2] == "）" and ("`" in t[:8] or "uics" in t[:20]):
        print("   " + t[:58])
print("\n=== 表2 全部行（结论列） ===")
for r in d2.tables[1].rows:
    print("   %-16s | %-10s | %s" % (r.cells[0].text.strip()[:16],
                                     r.cells[4].text.strip()[:10],
                                     r.cells[2].text.strip()[:54]))
print("\n段数=%d 表数=%d" % (len(d2.paragraphs), len(d2.tables)))
