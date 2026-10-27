# -*- coding: utf-8 -*-
import os, sys, glob, datetime
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

print("=" * 100)
print("一、根目录交付物（docx / pdf / pptx / png）")
print("=" * 100)
rows = []
for pat in ("*.docx", "*.pdf", "*.pptx", "*.png"):
    for f in glob.glob(os.path.join(r"E:\FPGA", pat)):
        st = os.stat(f)
        rows.append((st.st_mtime, st.st_size, os.path.basename(f)))
rows.sort(reverse=True)
for mt, sz, name in rows:
    print("  %s  %9d B  %s" % (datetime.datetime.fromtimestamp(mt).strftime("%m-%d %H:%M"), sz, name))

print()
print("=" * 100)
print("二、报告待补充项")
print("=" * 100)
from docx import Document
rep = r"E:\FPGA\比赛报告_填空版_20261023.docx"
d = Document(rep)
hit = 0
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if "待补充" in t or "待插入" in t:
        hit += 1
        print("  ¶%d  %s" % (i, t[:90]))
print("  合计 %d 处" % hit)

print()
print("=" * 100)
print("三、学习清单的代码块缺陷")
print("=" * 100)
d2 = Document(r"E:\FPGA\项目学习清单_20261023.docx")
ps = d2.paragraphs
blocks = []
i = 0
while i < len(ps):
    if ps[i].text.strip().startswith("\u258e"):
        st = i
        j = i + 1
        nbody = 0
        while j < len(ps):
            tt = ps[j].text.strip()
            if tt.startswith("\u25b8") or tt.startswith("\u258e"):
                break
            if ps[j].runs and ps[j].runs[0].font.size and ps[j].runs[0].font.size.pt >= 15:
                break
            nbody += 1
            j += 1
        has = j < len(ps) and ps[j].text.strip().startswith("\u25b8")
        blocks.append((ps[st].text.strip(), nbody, has))
        i = j
    else:
        i += 1
from collections import Counter
dup = [k for k, v in Counter(b[0] for b in blocks).items() if v > 1]
nonote = [b for b in blocks if not b[2]]
print("  代码块 %d 个；缺 ▸ 解释 %d 个；重复 %d 组" % (len(blocks), len(nonote), len(dup)))
for h, n, _ in nonote:
    print("    缺解释 [%d 行] %s" % (n, h))
for k in dup:
    print("    重复  x2 %s" % k)

print()
print("=" * 100)
print("四、其他交付物页数/段落数")
print("=" * 100)
for f in ("演示视频脚本_20261020.docx", "演示视频旁白稿_20261023.docx", "比赛报告_填空版_20261023.docx",
          "分镜检查表_20261023.docx", "心得体会写作指引_20261021.docx",
          "心得体会模板_20261019.docx", "给成员B_项目上手指南_20261020.docx",
          "给成员C_项目上手指南_20261020.docx", "项目学习清单_20261023.docx"):
    p = os.path.join(r"E:\FPGA", f)
    if os.path.exists(p):
        dd = Document(p)
        print("  %-38s %2d 段  %2d 表  %d B" % (f, len(dd.paragraphs), len(dd.tables), os.path.getsize(p)))
    else:
        print("  %-38s 不存在" % f)
